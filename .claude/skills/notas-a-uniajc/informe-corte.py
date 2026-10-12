# -*- coding: utf-8 -*-
"""Informe de notas de un corte con la marca de la UNIAJC (Word).

Marca: colores y tipografía de universidades/UNIAJ/config/universidades/uniajc.json y los
estilos de universidades/UNIAJ/config/parciales/parcial_docx_engine.py (banda azul #095292,
Calibri, gris #333333, sombreado alterno #F2F2F2) más el logo horizontal azul.

    python .claude/skills/notas-a-uniajc/informe-corte.py --corte 1 [--fecha "lunes 12 de octubre de 2026"]

Lee de la carpeta del corte (por omisión privado/notas-<periodo>/corte<N>/, ignorada por git):
  - los CSV de leer-libro.cjs (uno por curso de cursos.json);
  - observaciones.json (opcional): {"<documento>": "texto"} o {"<documento>": {"texto": "...",
    "registrada": "—"}} — "registrada" pisa lo que se mostraría (p. ej. «—» para quien no figura
    en el listado oficial y por eso no se registró). La clave «<archivo>|<documento>» acota la
    observación a un curso: un mismo estudiante puede estar en dos (la cohorte 341C);
  - ajustes.json (opcional): lista de frases para «Ajustes hechos en ExamLab antes de registrar».
Escribe «Informe de notas - Corte N - <periodo>.docx» en esa misma carpeta.
"""
import argparse
import csv
import datetime
import decimal
import json
import os
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, "..", "..", ".."))
CONFIG = json.load(open(os.path.join(AQUI, "cursos.json"), encoding="utf-8"))
LOGO = os.path.join(RAIZ, "universidades", "UNIAJ", "config", "slides", "assets", "logo_uniajc_hor_azul.png")
AZUL = RGBColor(0x09, 0x52, 0x92)
GRIS = RGBColor(0x33, 0x33, 0x33)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def n(s):
    return "".join(c for c in unicodedata.normalize("NFD", s or "") if unicodedata.category(c) != "Mn").upper().strip()


def un_decimal(x):
    """Como el portal: un decimal, mitad hacia arriba (4,47 → 4,5)."""
    return float(decimal.Decimal(str(x)).quantize(decimal.Decimal("0.1"), rounding=decimal.ROUND_HALF_UP))


def coma(x, dec=2):
    return f"{x:.{dec}f}".replace(".", ",")


def shade(el, fill):
    pr = el.get_or_add_pPr() if el.tag.endswith("}p") else el.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), fill)
    pr.append(s)


def run(p, text, size=11, bold=False, color=GRIS):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color
    return r


def para(doc, text="", size=11, bold=False, color=GRIS, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, before=0, fill=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if fill:
        shade(p._p, fill)
    if text:
        run(p, text, size, bold, color)
    return p


def banda(doc, text):
    return para(doc, "  " + text, 13, True, BLANCO, after=8, before=12, fill="095292")


def vinetas(doc, textos):
    for txt in textos:
        q = doc.add_paragraph(style="List Bullet")
        q.paragraph_format.space_after = Pt(2)
        run(q, txt, 10.5)


def encabezado(t, columnas, anchos, size=9.5):
    for i, c in enumerate(columnas):
        cel = t.rows[0].cells[i]
        cel.width = Inches(anchos[i])
        shade(cel._tc, "095292")
        run(cel.paragraphs[0], c, size, True, BLANCO)


def leer_json(ruta, por_omision):
    return json.load(open(ruta, encoding="utf-8")) if os.path.exists(ruta) else por_omision


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corte", type=int, default=1)
    ap.add_argument("--carpeta")
    ap.add_argument("--fecha", help="fecha del registro en el portal, en letras")
    a = ap.parse_args()
    corte = a.corte
    carpeta = a.carpeta or os.path.join(RAIZ, "privado", f"notas-{CONFIG['periodo']}", f"corte{corte}")
    hoy = datetime.date.today()
    fecha = a.fecha or f"{DIAS[hoy.weekday()]} {hoy.day} de {MESES[hoy.month - 1]} de {hoy.year}"
    evaluacion = CONFIG["evaluaciones"].get(str(corte), f"Corte {corte}")
    obs = leer_json(os.path.join(carpeta, "observaciones.json"), {})
    ajustes = leer_json(os.path.join(carpeta, "ajustes.json"), [])

    datos, resumen = [], []
    for curso in CONFIG["cursos"]:
        ruta = os.path.join(carpeta, f"Corte{corte} - {curso['archivo']}.csv")
        if not os.path.exists(ruta):
            raise SystemExit(f"Falta {ruta}: correr primero leer-libro.cjs --corte={corte}")
        with open(ruta, encoding="utf-8-sig") as f:
            filas = list(csv.reader(f, delimiter=";"))[1:]
        rows = []
        for est, documento, _correo, nota, *_ in filas:
            x = float(nota.replace(",", "."))
            registrada = "NP" if x <= 0 else coma(un_decimal(x), 1)
            # «<archivo>|<documento>» manda sobre «<documento>»: la cohorte 341C está en dos cursos.
            o = obs.get(f"{curso['archivo']}|{documento}", obs.get(documento, ""))
            texto = o.get("texto", "") if isinstance(o, dict) else o
            if isinstance(o, dict) and o.get("registrada"):
                registrada = o["registrada"]
            if x <= 0 and not texto and registrada == "NP":
                texto = "No Presentó: sin actividad registrada en el corte."
            rows.append((est, documento, x, registrada, texto))
        oficiales = [r for r in rows if r[3] != "—"]
        con_nota = [un_decimal(r[2]) for r in oficiales if r[3] != "NP"]
        aprueban = sum(1 for v in con_nota if v >= 3.0)
        datos.append((curso, rows))
        resumen.append((f"{curso['asignatura']} ({curso['grupo']})", len(oficiales),
                        coma(sum(r[2] for r in oficiales) / len(oficiales)) if oficiales else "—",
                        aprueban, len(con_nota) - aprueban, sum(1 for r in oficiales if r[3] == "NP")))

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(0.7), Inches(0.7)
        s.left_margin, s.right_margin = Inches(0.75), Inches(0.75)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(11)

    if os.path.exists(LOGO):
        para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4).add_run().add_picture(LOGO, width=Inches(2.4))
    para(doc, "Institución Universitaria Antonio José Camacho — UNIAJC", 10, True, AZUL, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "Facultad de Ingenierías · Programa de Ingeniería de Sistemas", 10, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, after=10)
    banda(doc, f"Informe de notas — Corte {corte} · Periodo {CONFIG['periodo']}")
    para(doc, f"Docente: {CONFIG['docente']}  ·  {CONFIG['correoDocente']}", 10, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, f"Registro en el sistema académico: {fecha}  ·  {evaluacion.capitalize() if evaluacion.isupper() else evaluacion}", 10, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, after=10)

    banda(doc, "1. Resumen")
    para(doc, f"Las notas del Corte {corte} quedaron registradas en el sistema académico de la universidad y se "
              "verificaron una por una: después de registrar, cada grupo se abrió de nuevo y se comparó contra la "
              "nota de ExamLab.", after=6)
    anchos = [3.0, 0.85, 0.8, 1.05, 0.85, 0.5]
    t = doc.add_table(rows=1, cols=6)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    encabezado(t, ["Asignatura (grupo)", "Estudiantes", "Promedio", "Aprueban (≥ 3,0)", "Reprueban", "NP"], anchos)
    tot = [0, 0, 0, 0]
    for k, r in enumerate(resumen):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].width = Inches(anchos[i])
            if k % 2 == 1:
                shade(cells[i]._tc, "F2F2F2")
            q = cells[i].paragraphs[0]
            if i:
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(q, str(v), 9.5)
        tot = [tot[0] + r[1], tot[1] + r[3], tot[2] + r[4], tot[3] + r[5]]
    cells = t.add_row().cells
    for i, v in enumerate(["Total", tot[0], "", tot[1], tot[2], tot[3]]):
        cells[i].width = Inches(anchos[i])
        shade(cells[i]._tc, "E8F4FA")
        q = cells[i].paragraphs[0]
        if i:
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(q, str(v), 9.5, True, AZUL)
    para(doc, "Promedio: nota del corte en ExamLab, con dos decimales. Aprueban y reprueban: nota registrada en la "
              "universidad, con un decimal. NP: No Presentó.", 9, False, GRIS, after=4, before=4)

    banda(doc, "2. Criterios del registro")
    vinetas(doc, [
        "La nota del corte es la que calcula ExamLab con los pesos del curso.",
        "El sistema de la universidad recibe la nota con un decimal: se redondea con la mitad hacia arriba (4,47 queda 4,5).",
        "NP (No Presentó) se registra a quien no tiene ninguna actividad en el corte; el sistema la toma como 0,0.",
        "Quien no tiene ninguna asistencia en el corte queda con 0. Si tiene trabajo propio en línea (una evaluación o "
        "un taller que entregó él), conserva esa nota: la asistencia ya le cuenta 0 dentro del cálculo.",
    ])

    if ajustes:
        banda(doc, "3. Ajustes hechos en ExamLab antes de registrar")
        vinetas(doc, ajustes)

    banda(doc, f"{4 if ajustes else 3}. Notas por grupo")
    anchos = [0.35, 2.35, 1.0, 0.75, 0.8, 2.05]
    for curso, rows in datos:
        para(doc, f"{curso['asignatura']} — Grupo {curso['grupo']} ({curso['codigo']})", 12, True, AZUL, before=10, after=4)
        t = doc.add_table(rows=1, cols=6)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        encabezado(t, ["N.º", "Estudiante", "Documento", "ExamLab", "Registrada", "Observación"], anchos, 9)
        for k, (est, documento, x, registrada, texto) in enumerate(sorted(rows, key=lambda r: n(r[0]))):
            cells = t.add_row().cells
            for i, v in enumerate([k + 1, est, documento, coma(x), registrada, texto]):
                cells[i].width = Inches(anchos[i])
                if texto:
                    shade(cells[i]._tc, "FFF4C2")
                elif k % 2 == 1:
                    shade(cells[i]._tc, "F2F2F2")
                q = cells[i].paragraphs[0]
                q.paragraph_format.space_after = Pt(0)
                if i in (0, 3, 4):
                    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(q, str(v), 8.5 if i == 5 else 9, i == 4)
        para(doc, "", after=4)

    banda(doc, f"{5 if ajustes else 4}. Para tener en cuenta")
    vinetas(doc, [
        "Las filas resaltadas en amarillo tienen una observación: un NP, un ajuste o un reclamo atendido.",
        "Los grupos no se bloquearon en el sistema: si hace falta corregir una nota, todavía se puede modificar.",
    ])

    salida = os.path.join(carpeta, f"Informe de notas - Corte {corte} - {CONFIG['periodo']}.docx")
    doc.save(salida)
    print("guardado:", os.path.relpath(salida, RAIZ))
    print("total", {"estudiantes": tot[0], "aprueban": tot[1], "reprueban": tot[2], "np": tot[3]})


if __name__ == "__main__":
    main()
